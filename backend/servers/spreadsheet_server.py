"""Spreadsheet MCP server. Reads CSV and Excel after the file type is known."""

from __future__ import annotations

import csv
from pathlib import Path

from servers.file_server import list_files
from servers.workspace import allowed_file, safe_path

NAME = "spreadsheet"
TOOLS = ["spreadsheet.list", "spreadsheet.read"]


def list_sheets() -> list[str]:
    return [item["path"] for item in list_files() if item["path"].lower().endswith(".csv")]


def read_sheet(relative_path: str, limit: int = 20) -> dict:
    path = safe_path(relative_path)
    return read_path(path, limit)


def read_path(path: Path, limit: int = 20) -> dict:
    target = allowed_file(path)
    suffix = target.suffix.lower()
    if suffix == ".csv":
        rows = _csv_rows(target)
    elif suffix == ".xlsx":
        rows = _xlsx_rows(target)
    else:
        raise ValueError(f"Unsupported spreadsheet type: {suffix or 'unknown'}")
    header = rows[0] if rows else []
    body = rows[1 : limit + 1]
    return {"path": str(target), "columns": header, "rows": body, "row_count": max(0, len(rows) - 1)}


def _csv_rows(path: Path) -> list[list[str]]:
    with path.open(encoding="utf-8", errors="ignore", newline="") as handle:
        return list(csv.reader(handle))


def _xlsx_rows(path: Path) -> list[list[str]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = []
    for index, row in enumerate(sheet.iter_rows(values_only=True)):
        if index >= 21:
            break
        rows.append(["" if cell is None else str(cell) for cell in row])
    workbook.close()
    return rows
